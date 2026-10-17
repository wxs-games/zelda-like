# -*- coding: utf-8 -*-
"""Update schedule xlsx: 17 new clips -> 已完成, real cell sizes, fixed sheet paths."""
from pathlib import Path
from PIL import Image
import openpyxl

# 脚本位于 <工程根>\Tools\art-pipeline\，向上两级定位工程根（含 Assets 的目录），免配置路径
art = Path(__file__).resolve().parents[2] / "Assets" / "Art"

# 1) verify real frame sizes per role
def frame_size(rel, name):
    p = art / rel / "Frames" / name / f"{name}_00.png"
    return "%d×%d" % Image.open(p).size

sizes = {
    "player": frame_size("Art/Player", "Player_Dash"),
    "enemy03": frame_size("Art/Enemies/Enemy03", "Enemy03_Spawn"),
    "tower01": frame_size("Art/Towers/Tower01", "Tower01_Aim"),
    "boss01": frame_size("Art/Boss/Boss01", "Boss01_Special"),
}
print("real cell sizes:", sizes)

xlsx = art / "2D游戏动画资源排期表.xlsx"
wb = openpyxl.load_workbook(xlsx)

# 2) player sheet: rows 11-23 (clips 10-22)
ws = wb["主角动作清单"]
for r in range(11, 24):
    ws.cell(row=r, column=14).value = sizes["player"]   # N 画布尺寸
    ws.cell(row=r, column=15).value = "已完成"          # O 制作状态

# 3) enemies & boss sheet: rows 12, 19, 26, 28
ws = wb["敌人与BOSS"]
fix = {
    12: (sizes["enemy03"], "Assets/Art/Enemies/Enemy03/Sheets/Enemy03_Spawn_sheet.png"),
    19: (sizes["tower01"], "Assets/Art/Towers/Tower01/Sheets/Tower01_Aim_sheet.png"),
    26: (sizes["boss01"],  "Assets/Art/Boss/Boss01/Sheets/Boss01_Special_sheet.png"),
    28: (sizes["boss01"],  "Assets/Art/Boss/Boss01/Sheets/Boss01_Enrage_sheet.png"),
}
for r, (size, path) in fix.items():
    ws.cell(row=r, column=12).value = path              # L 精灵表路径
    ws.cell(row=r, column=14).value = size              # N 画布尺寸
    ws.cell(row=r, column=15).value = "已完成"          # O 制作状态

wb.save(xlsx)
print("[OK] xlsx updated:", xlsx.name)

# 4) verify no 未开始 remains
wb2 = openpyxl.load_workbook(xlsx)
for name in ["主角动作清单", "敌人与BOSS"]:
    ws2 = wb2[name]
    left = [r for r in range(2, ws2.max_row + 1) if ws2.cell(row=r, column=15).value == "未开始"]
    print(name, "未开始 rows:", left if left else "none")
